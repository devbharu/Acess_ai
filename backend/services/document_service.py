import io
import os
import re
import logging
from pathlib import Path
from PIL import Image
import pytesseract
from typing import List, Dict, Any, Optional

try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    import docx
except ImportError:
    docx = None

logger = logging.getLogger(__name__)

class DocumentService:
    @staticmethod
    def extract_from_file(file_path: str, filename: str) -> Dict[str, Any]:
        """
        Extract text from PDF, DOCX, TXT, MD, etc.
        Applies Tesseract OCR if document is scanned or text is sparse.
        Returns:
            pages (int), chars (int), ocr_used (bool), pages_data (list of dicts {page, text})
        """
        ext = Path(filename).suffix.lower()
        if ext == ".pdf":
            return DocumentService._extract_pdf(file_path)
        elif ext in [".docx", ".doc"]:
            return DocumentService._extract_docx(file_path)
        else:
            return DocumentService._extract_plain_text(file_path)

    @staticmethod
    def _extract_pdf(file_path: str) -> Dict[str, Any]:
        pages_data = []
        ocr_used = False
        total_chars = 0

        if fitz:
            doc = fitz.open(file_path)
            total_pages = len(doc)
            for page_num in range(total_pages):
                page = doc[page_num]
                text = page.get_text().strip()

                # If text is very short (< 30 chars), perform OCR on rendered page image
                if len(text) < 30:
                    try:
                        pix = page.get_pixmap(dpi=150)
                        img = Image.open(io.BytesIO(pix.tobytes("png")))
                        ocr_text = pytesseract.image_to_string(img).strip()
                        if ocr_text:
                            text = ocr_text
                            ocr_used = True
                    except Exception as e:
                        logger.warning("OCR failed on PDF page %d: %s", page_num + 1, e)

                pages_data.append({"page": page_num + 1, "text": text})
                total_chars += len(text)
            doc.close()
        else:
            with open(file_path, "rb") as f:
                content = f.read().decode("utf-8", errors="ignore")
            pages_data.append({"page": 1, "text": content})
            total_chars = len(content)

        return {
            "pages": len(pages_data),
            "chars": total_chars,
            "ocr_used": ocr_used,
            "pages_data": pages_data,
        }

    @staticmethod
    def _extract_docx(file_path: str) -> Dict[str, Any]:
        if docx:
            doc = docx.Document(file_path)
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text.strip())
            for table in doc.tables:
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_text:
                        full_text.append(" | ".join(row_text))
            text = "\n\n".join(full_text)
        else:
            with open(file_path, "rb") as f:
                text = f.read().decode("utf-8", errors="ignore")

        return {
            "pages": 1,
            "chars": len(text),
            "ocr_used": False,
            "pages_data": [{"page": 1, "text": text}],
        }

    @staticmethod
    def _extract_plain_text(file_path: str) -> Dict[str, Any]:
        with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
            text = f.read()

        return {
            "pages": 1,
            "chars": len(text),
            "ocr_used": False,
            "pages_data": [{"page": 1, "text": text}],
        }

    @staticmethod
    def get_document_outline(file_path: str) -> List[Dict[str, Any]]:
        """
        Extracts document structure: TOC outline or heading blocks per page.
        """
        ext = Path(file_path).suffix.lower()
        if ext != ".pdf" or not fitz:
            return [{"page": 1, "title": Path(file_path).name, "level": 1}]

        outline = []
        try:
            doc = fitz.open(file_path)
            toc = doc.get_toc()
            # If doc has native TOC with descriptive titles
            if toc and any(item[1] and not item[1].lower().startswith("slide") for item in toc):
                for item in toc:
                    outline.append({
                        "level": item[0],
                        "title": item[1].strip(),
                        "page": max(1, item[2]),
                    })
            else:
                # Heuristically extract prominent headings from blocks on each page
                for i, page in enumerate(doc):
                    blocks = page.get_text("blocks")
                    title = ""
                    for b in blocks:
                        b_text = b[4].strip()
                        if len(b_text) > 3:
                            # Clean newline
                            clean_line = " ".join(b_text.splitlines())
                            title = clean_line[:80]
                            break
                    if not title:
                        title = f"Page {i + 1}"
                    outline.append({
                        "level": 1,
                        "title": title,
                        "page": i + 1,
                    })
            doc.close()
        except Exception as e:
            logger.warning("Error getting outline for %s: %s", file_path, e)
            outline = [{"page": 1, "title": Path(file_path).name, "level": 1}]

        return outline

    @staticmethod
    def get_page_content(file_path: str, page_number: int) -> Dict[str, Any]:
        """
        Retrieves text, OCR, and metadata for a specific 1-indexed page.
        """
        ext = Path(file_path).suffix.lower()
        if ext != ".pdf" or not fitz:
            ext_data = DocumentService.extract_from_file(file_path, Path(file_path).name)
            return {
                "page": 1,
                "text": ext_data["pages_data"][0]["text"] if ext_data["pages_data"] else "",
                "ocr_used": ext_data["ocr_used"],
                "images_count": 0,
            }

        try:
            doc = fitz.open(file_path)
            idx = max(0, min(page_number - 1, len(doc) - 1))
            page = doc[idx]
            text = page.get_text().strip()
            ocr_used = False

            if len(text) < 30:
                try:
                    pix = page.get_pixmap(dpi=150)
                    img = Image.open(io.BytesIO(pix.tobytes("png")))
                    ocr = pytesseract.image_to_string(img).strip()
                    if ocr:
                        text = ocr
                        ocr_used = True
                except Exception:
                    pass

            images_count = len(page.get_images())
            doc.close()
            return {
                "page": idx + 1,
                "text": text,
                "ocr_used": ocr_used,
                "images_count": images_count,
            }
        except Exception as e:
            logger.error("Error reading page %d from %s: %s", page_number, file_path, e)
            return {"page": page_number, "text": "", "ocr_used": False, "images_count": 0, "error": str(e)}

    @staticmethod
    def search_document(file_path: str, query: str, max_matches: int = 6) -> List[Dict[str, Any]]:
        """
        Searches extracted text across all pages for a query phrase or keywords.
        Returns matching pages, exact snippet, and context.
        """
        if not query or not query.strip():
            return []

        clean_q = query.strip().lower()
        q_words = [w for w in re.findall(r"\w+", clean_q) if len(w) > 2]
        matches = []

        ext_data = DocumentService.extract_from_file(file_path, Path(file_path).name)
        for p in ext_data["pages_data"]:
            page_text = p["text"]
            lower_text = page_text.lower()

            score = 0
            # Direct phrase match
            if clean_q in lower_text:
                score += 10
            # Word overlap
            matched_words = [w for w in q_words if w in lower_text]
            score += len(matched_words)

            if score > 0:
                # Find best snippet window
                idx = -1
                if clean_q in lower_text:
                    idx = lower_text.find(clean_q)
                elif matched_words:
                    idx = lower_text.find(matched_words[0])

                if idx >= 0:
                    start = max(0, idx - 120)
                    end = min(len(page_text), idx + 280)
                    snippet = page_text[start:end].strip().replace("\n", " ")
                    if start > 0:
                        snippet = "..." + snippet
                    if end < len(page_text):
                        snippet = snippet + "..."
                else:
                    snippet = page_text[:250].strip().replace("\n", " ") + "..."

                matches.append({
                    "page": p["page"],
                    "score": score,
                    "snippet": snippet,
                    "matched_words": matched_words,
                })

        matches.sort(key=lambda x: x["score"], reverse=True)
        return matches[:max_matches]

    @staticmethod
    def extract_tables(file_path: str, page_number: Optional[int] = None) -> List[Dict[str, Any]]:
        """
        Extracts structured tables from the document or specific page.
        """
        ext = Path(file_path).suffix.lower()
        tables_res = []

        if ext == ".pdf" and fitz:
            try:
                doc = fitz.open(file_path)
                pages_to_check = [page_number - 1] if page_number is not None else range(len(doc))
                for p_idx in pages_to_check:
                    if 0 <= p_idx < len(doc):
                        page = doc[p_idx]
                        try:
                            tabs = page.find_tables()
                            for t_idx, tab in enumerate(tabs.tables):
                                df = tab.extract()
                                if df:
                                    # Convert to markdown table format
                                    headers = [str(c or "") for c in df[0]]
                                    md_rows = ["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"]
                                    for row in df[1:]:
                                        clean_row = [str(c or "").replace("\n", " ") for c in row]
                                        md_rows.append("| " + " | ".join(clean_row) + " |")
                                    tables_res.append({
                                        "page": p_idx + 1,
                                        "table_index": t_idx + 1,
                                        "rows": len(df),
                                        "cols": len(headers),
                                        "markdown": "\n".join(md_rows),
                                        "raw_data": df,
                                    })
                        except Exception:
                            pass
                doc.close()
            except Exception as e:
                logger.warning("Table extraction error: %s", e)

        # Fallback table search using pattern in plain text
        if not tables_res:
            ext_data = DocumentService.extract_from_file(file_path, Path(file_path).name)
            for p in ext_data["pages_data"]:
                lines = [l for l in p["text"].splitlines() if "|" in l or "\t" in l]
                if len(lines) >= 2:
                    tables_res.append({
                        "page": p["page"],
                        "table_index": 1,
                        "rows": len(lines),
                        "cols": 2,
                        "markdown": "\n".join(lines[:10]),
                        "raw_data": [l.split("|") for l in lines[:10]],
                    })

        return tables_res

    @staticmethod
    def get_document_images(file_path: str, page_number: Optional[int] = None) -> List[Dict[str, Any]]:
        """
        Extracts metadata of embedded images in PDF.
        """
        ext = Path(file_path).suffix.lower()
        images_info = []
        if ext == ".pdf" and fitz:
            try:
                doc = fitz.open(file_path)
                pages_to_check = [page_number - 1] if page_number is not None else range(len(doc))
                for p_idx in pages_to_check:
                    if 0 <= p_idx < len(doc):
                        page = doc[p_idx]
                        for img_idx, img in enumerate(page.get_images()):
                            xref = img[0]
                            base_img = doc.extract_image(xref)
                            images_info.append({
                                "page": p_idx + 1,
                                "image_index": img_idx + 1,
                                "width": base_img.get("width"),
                                "height": base_img.get("height"),
                                "extension": base_img.get("ext"),
                                "size_bytes": len(base_img.get("image", b"")),
                            })
                doc.close()
            except Exception as e:
                logger.warning("Error getting document images: %s", e)
        return images_info
